"""Create the downloadable Word edition from the same verified daily JSON."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from render import validate


def add_source_link(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '006D70'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def save_docx(data, path):
    day, total = validate(data)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [('Normal', 11), ('Title', 24), ('Heading 1', 15), ('Subtitle', 10)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.line_spacing = 1.2
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
    doc.add_paragraph('每日新闻简报', 'Title')
    doc.add_paragraph(f'{day.isoformat()}  纽约日期  共 {total} 条精选', 'Subtitle')
    doc.add_paragraph('本简报汇集生成前24小时内筛选的新闻。外语报道以中文简要译述，点击来源可阅读原始报道。')
    for block in data['sections']:
        doc.add_heading(block['name'], level=1)
        for item in block['items']:
            summary = doc.add_paragraph(item['summary'])
            summary.paragraph_format.keep_with_next = True
            summary.paragraph_format.keep_together = True
            source = doc.add_paragraph()
            source.paragraph_format.space_after = Pt(10)
            source.paragraph_format.keep_together = True
            add_source_link(source, item['source'] + ' 阅读原文', item['url'])
        if block.get('note'):
            doc.add_paragraph(block['note'])
    doc.add_paragraph('概要由 AI 根据所列来源整理，详情以原文为准。')
    footer = section.footer.paragraphs[0]
    footer.add_run('每日新闻简报  ' + day.isoformat() + '   第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    for run in footer.runs: run.font.size = Pt(9)
    doc.core_properties.title = '每日新闻简报 ' + day.isoformat()
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.subject = '中文概要与原始信源链接'
    doc.save(path)
