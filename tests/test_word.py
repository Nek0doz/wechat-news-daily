import io
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from export_word import save_docx
from render import render
from test_digest import fixture


class WordTests(unittest.TestCase):
    def test_word_preserves_all_sections_summaries_and_original_urls(self):
        data = fixture()
        data['sections'][2]['items'] = []
        data['sections'][2]['note'] = '暂无可靠的重要更新。'
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / '2026-09-16.docx'
            save_docx(data, path)
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            urls = {r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK}
            expected = {i['url'] for b in data['sections'] for i in b['items']}
            self.assertEqual(urls, expected)
            for block in data['sections']:
                self.assertIn(block['name'], text)
                for item in block['items']: self.assertIn(item['summary'], text)
                if block.get('note'): self.assertIn(block['note'], text)
            self.assertIn('href="2026-09-16.docx" download', render(data))

    def test_weekly_word_preserves_56_items_and_links(self):
        data = fixture(); data['period'] = 'weekly'
        for index, block in enumerate(data['sections']):
            template = block['items'][0]
            block['items'] = [dict(template, summary=f'第{index}板块第{j}条周报概要',
                                   url=f'https://example.org/weekly/{index}/{j}',
                                   published_at='2026-09-10T09:00:00-04:00') for j in range(8)]
        buffer = io.BytesIO()
        save_docx(data, buffer)
        doc = Document(io.BytesIO(buffer.getvalue()))
        text = '\\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('每周新闻简报', text)
        self.assertIn('生成前7天', text)
        self.assertIn('共 56 条精选', text)
        for block in data['sections']:
            self.assertIn(block['name'], text)
            for item in block['items']:
                self.assertIn(item['summary'], text)
        urls = {r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK}
        self.assertEqual(urls, {i['url'] for b in data['sections'] for i in b['items']})


if __name__ == '__main__': unittest.main()
